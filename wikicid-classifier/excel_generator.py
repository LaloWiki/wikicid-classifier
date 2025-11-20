"""
Generador de Excel con empresas clasificadas por sector
ESTRUCTURA: Cada sector en una pestaña, agrupado por tipo de IA
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import logging
from typing import List, Dict
from pathlib import Path

logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class ExcelGenerator:
    def __init__(self, output_path: str):
        self.output_path = output_path
        
        # Colores para headers
        self.header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        self.header_font = Font(bold=True, color="FFFFFF", size=11)
        
        # Bordes
        self.thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )
        
        # Sectores disponibles
        self.sectores = [
            'Sector financiero',
            'Telecomunicaciones',
            'Retail eCommerce',
            'Salud'
        ]
        
        # Orden de tipos de IA
        self.tipos_ia_orden = [
            'IA tradicional ML clásico',
            'IA Generativa (Gen IA)',
            'IA Agéntica'
        ]
    
    def generar_excel(self, empresas_clasificadas: List[Dict]):
        """
        Genera el Excel con pestañas por sector
        """
        logger.info("📊 Generando Excel con empresas clasificadas...")
        
        # Validar que hay datos
        if not empresas_clasificadas:
            logger.warning("⚠️  No hay empresas para procesar")
            return None
        
        try:
            # Crear writer de Excel
            with pd.ExcelWriter(self.output_path, engine='openpyxl') as writer:
                
                # Generar una pestaña por sector
                for sector in self.sectores:
                    df_sector = self._crear_dataframe_sector(empresas_clasificadas, sector)
                    
                    if not df_sector.empty:
                        # Escribir al Excel
                        sheet_name = self._get_sheet_name(sector)
                        df_sector.to_excel(writer, sheet_name=sheet_name, index=False)
                        
                        # Aplicar formato
                        self._aplicar_formato(writer.book[sheet_name])
                        
                        logger.info(f"✓ Pestaña '{sheet_name}' creada con {len(df_sector)} registros")
                    else:
                        logger.info(f"ℹ️  Sector '{sector}' sin empresas - pestaña no creada")
                
                # Crear pestaña de resumen
                df_resumen = self._crear_resumen(empresas_clasificadas)
                df_resumen.to_excel(writer, sheet_name='Resumen', index=False)
                self._aplicar_formato_resumen(writer.book['Resumen'])
                
                logger.info("✓ Pestaña 'Resumen' creada")
            
            logger.info(f"✅ Excel generado exitosamente: {self.output_path}")
            return self.output_path
            
        except Exception as e:
            logger.error(f"❌ Error al generar Excel: {str(e)}")
            raise
    
    def _crear_dataframe_sector(self, empresas: List[Dict], sector: str) -> pd.DataFrame:
        """
        Crea un DataFrame con las empresas de un sector específico
        ESTRUCTURA: Agrupado por tipo de IA, con casos de uso específicos
        """
        filas = []
        
        for empresa in empresas:
            clasificacion = empresa.get('clasificacion', {})
            sectores = clasificacion.get('sectores', [])
            
            # Si la empresa pertenece a este sector
            if sector in sectores:
                tipo_ia = clasificacion.get('tipo_ia', 'Sin clasificar')
                casos_uso_sector = clasificacion.get('casos_uso', {}).get(sector, [])
                descripcion = clasificacion.get('descripcion_relevante', '')
                nivel_ajuste = clasificacion.get('nivel_ajuste', '')
                observaciones = clasificacion.get('observaciones', '')
                
                # Crear una fila por cada caso de uso
                if casos_uso_sector:
                    for idx, caso_uso in enumerate(casos_uso_sector):
                        fila = {
                            'TIPO DE IA': tipo_ia,
                            'Casos de uso/ Vertical': caso_uso,
                            'Empresa': empresa.get('nombre', 'Sin nombre'),
                            'Descripción relevante': descripcion if idx == 0 else '',
                            'Nivel de ajuste': nivel_ajuste if idx == 0 else '',
                            'Observaciones': observaciones if idx == 0 else ''
                        }
                        filas.append(fila)
                else:
                    # Si no tiene casos de uso, crear fila genérica
                    fila = {
                        'TIPO DE IA': tipo_ia,
                        'Casos de uso/ Vertical': 'Sin caso de uso específico',
                        'Empresa': empresa.get('nombre', 'Sin nombre'),
                        'Descripción relevante': descripcion,
                        'Nivel de ajuste': nivel_ajuste,
                        'Observaciones': observaciones
                    }
                    filas.append(fila)
        
        df = pd.DataFrame(filas)
        
        # Ordenar por tipo de IA (según orden definido) y luego por caso de uso
        if not df.empty:
            # Crear columna auxiliar para ordenar por tipo de IA
            tipo_ia_map = {tipo: idx for idx, tipo in enumerate(self.tipos_ia_orden)}
            df['_orden_tipo'] = df['TIPO DE IA'].map(lambda x: tipo_ia_map.get(x, 999))
            df = df.sort_values(['_orden_tipo', 'Casos de uso/ Vertical', 'Empresa'])
            df = df.drop('_orden_tipo', axis=1)
        
        return df
    
    def _crear_resumen(self, empresas: List[Dict]) -> pd.DataFrame:
        """
        Crea un DataFrame de resumen con estadísticas
        """
        resumen_data = []
        
        # Estadísticas generales
        resumen_data.append({
            'Métrica': 'Total de empresas procesadas',
            'Valor': len(empresas)
        })
        
        clasificadas = [e for e in empresas if e.get('clasificacion')]
        resumen_data.append({
            'Métrica': 'Empresas clasificadas exitosamente',
            'Valor': len(clasificadas)
        })
        
        resumen_data.append({
            'Métrica': 'Empresas sin clasificar',
            'Valor': len(empresas) - len(clasificadas)
        })
        
        resumen_data.append({'Métrica': '', 'Valor': ''})  # Separador
        
        # Contar por sector
        for sector in self.sectores:
            count = sum(1 for e in empresas 
                       if sector in e.get('clasificacion', {}).get('sectores', []))
            resumen_data.append({
                'Métrica': f'Empresas en {sector}',
                'Valor': count
            })
        
        resumen_data.append({'Métrica': '', 'Valor': ''})  # Separador
        
        # Contar por tipo de IA
        for tipo_ia in self.tipos_ia_orden:
            count = sum(1 for e in empresas 
                       if e.get('clasificacion', {}).get('tipo_ia') == tipo_ia)
            resumen_data.append({
                'Métrica': f'Empresas con {tipo_ia}',
                'Valor': count
            })
        
        df = pd.DataFrame(resumen_data)
        return df
    
    def _get_sheet_name(self, sector: str) -> str:
        """
        Convierte el nombre del sector a un nombre válido de hoja
        (Excel limita a 31 caracteres)
        """
        nombres = {
            'Sector financiero': 'Financiero',
            'Telecomunicaciones': 'Telecomunicaciones',
            'Retail eCommerce': 'Retail eCommerce',
            'Salud': 'Salud'
        }
        nombre = nombres.get(sector, sector)
        return nombre[:31]  # Límite de Excel
    
    def _aplicar_formato(self, worksheet):
        """
        Aplica formato a una hoja de Excel con datos de empresas
        """
        # Formatear encabezados
        for cell in worksheet[1]:
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = self.thin_border
        
        # Ajustar anchos de columna
        column_widths = {
            'A': 30,  # TIPO DE IA
            'B': 55,  # Casos de uso
            'C': 30,  # Empresa
            'D': 70,  # Descripción
            'E': 15,  # Nivel
            'F': 40   # Observaciones
        }
        
        for col, width in column_widths.items():
            worksheet.column_dimensions[col].width = width
        
        # Aplicar bordes y alineación a todas las celdas
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for cell in row:
                cell.border = self.thin_border
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        
        # Congelar primera fila
        worksheet.freeze_panes = 'A2'
        
        # Ajustar altura de filas para mejor legibilidad
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            worksheet.row_dimensions[row[0].row].height = 35
    
    def _aplicar_formato_resumen(self, worksheet):
        """
        Aplica formato especial a la hoja de resumen
        """
        # Formatear encabezados
        for cell in worksheet[1]:
            cell.fill = PatternFill(start_color="70AD47", end_color="70AD47", fill_type="solid")
            cell.font = Font(bold=True, color="FFFFFF", size=12)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = self.thin_border
        
        # Ajustar anchos
        worksheet.column_dimensions['A'].width = 50
        worksheet.column_dimensions['B'].width = 20
        
        # Formatear celdas de datos
        for row in worksheet.iter_rows(min_row=2, max_row=worksheet.max_row):
            for idx, cell in enumerate(row):
                cell.border = self.thin_border
                if idx == 0:  # Columna de métrica
                    cell.alignment = Alignment(horizontal='left', vertical='center')
                    if cell.value == '':  # Separadores
                        cell.fill = PatternFill(start_color="E7E6E6", end_color="E7E6E6", fill_type="solid")
                else:  # Columna de valor
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                    cell.font = Font(bold=True, size=11)
        
        worksheet.freeze_panes = 'A2'


# Función auxiliar para uso fácil
def generar_excel_empresas(empresas: List[Dict], output_path: str = "empresas_ia_clasificadas.xlsx"):
    """
    Función wrapper para generar el Excel fácilmente
    
    Args:
        empresas: Lista de diccionarios con la estructura:
            {
                'nombre': str,
                'clasificacion': {
                    'tipo_ia': str,
                    'sectores': List[str],
                    'casos_uso': Dict[str, List[str]],
                    'descripcion_relevante': str,
                    'nivel_ajuste': str,
                    'observaciones': str
                }
            }
        output_path: Ruta del archivo Excel a generar
    
    Returns:
        str: Ruta del archivo generado
    """
    generator = ExcelGenerator(output_path)
    return generator.generar_excel(empresas)


# Ejemplo de uso
if __name__ == "__main__":
    # Datos de ejemplo
    empresas_ejemplo = [
        {
            'nombre': 'TechBank AI',
            'clasificacion': {
                'tipo_ia': 'IA Generativa (Gen IA)',
                'sectores': ['Sector financiero'],
                'casos_uso': {
                    'Sector financiero': [
                        'Análisis de riesgo crediticio',
                        'Detección de fraude'
                    ]
                },
                'descripcion_relevante': 'Plataforma de IA para análisis financiero',
                'nivel_ajuste': 'Alto',
                'observaciones': 'Implementación exitosa en 5 bancos'
            }
        },
        {
            'nombre': 'HealthAI Solutions',
            'clasificacion': {
                'tipo_ia': 'IA tradicional ML clásico',
                'sectores': ['Salud'],
                'casos_uso': {
                    'Salud': ['Diagnóstico asistido', 'Predicción de enfermedades']
                },
                'descripcion_relevante': 'Sistema de diagnóstico médico con ML',
                'nivel_ajuste': 'Medio',
                'observaciones': 'En fase de certificación'
            }
        }
    ]
    
    # Generar Excel
    resultado = generar_excel_empresas(empresas_ejemplo, "output_ejemplo.xlsx")
    print(f"Excel generado: {resultado}")
